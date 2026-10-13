"""Generate an editable research review deck from committed model results."""
from pathlib import Path
import os,sys,json,math
os.environ.setdefault('MPLCONFIGDIR','/tmp/hx-presentation-mpl')
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from roughness import build_roughness
OUT=ROOT/'presentations';ASSETS=OUT/'assets';ASSETS.mkdir(parents=True,exist_ok=True)
smooth=pd.read_csv(ROOT/'results/smooth_baseline/heldout_predictions.csv')
dimple=pd.read_csv(ROOT/'results/smooth_baseline/Dimpled-transfer.summary.csv')
old=pd.read_csv(ROOT/'results/Spencer-vs-Numerical-HTC-All-40-Cases.csv')
audit=pd.read_csv(ROOT/'validation/Correlation_audit_first_pass_HEOS.csv')
fine=pd.read_csv(ROOT/'results/smooth_baseline/Dimpled_fine_grid_checks.csv')
model=json.loads((ROOT/'results/smooth_baseline/model.json').read_text())
NAVY='142B45';TEAL='008A83';BLUE='327EC0';ORANGE='D88132';RED='BE4F50';GRAY='536478';LIGHT='F1F5F8'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':'#'+NAVY,'text.color':'#'+NAVY,'figure.facecolor':'white','axes.facecolor':'white','savefig.dpi':190})
def save(fig,name):
 fig.savefig(ASSETS/name,bbox_inches='tight',facecolor='white');plt.close(fig);return ASSETS/name
fig,ax=plt.subplots(figsize=(8.8,4.4))
for col,label,c in [('uncorrected_htc_W_m2K','Uncorrected DH',BLUE),('heldout_numerical_htc_W_m2K','Smooth correction: held-out',TEAL)]:
 ax.scatter(smooth.target_htc_W_m2K,smooth[col],s=38,color='#'+c,label=label,alpha=.8)
v=np.linspace(0,2600);ax.plot(v,v,color='#'+GRAY,lw=1,label='Exact agreement');ax.fill_between(v,.95*v,1.05*v,color='#'+GRAY,alpha=.10,label='±5% target band')
ax.set(xlabel='Spencer target HTC (W/m²·K)',ylabel='Numerical HTC (W/m²·K)',xlim=(0,2600),ylim=(0,2600));ax.legend(fontsize=9,loc='upper left');save(fig,'smooth_parity.png')
fig,ax=plt.subplots(figsize=(10.5,4.1));x=np.arange(len(smooth));ax.axhspan(-5,5,color='#'+TEAL,alpha=.1)
ax.plot(x,smooth.uncorrected_error_percent,'o-',color='#'+BLUE,ms=4,lw=1,label='Uncorrected DH');ax.plot(x,smooth.heldout_error_percent,'o-',color='#'+TEAL,ms=4,lw=1,label='Held-out correction')
ax.axhline(0,color='#'+GRAY,lw=.7);ax.set_xticks(x);ax.set_xticklabels(smooth.Case,rotation=90,fontsize=7);ax.set_ylabel('Signed HTC error (%)');ax.legend(fontsize=9,ncol=2);save(fig,'smooth_case_errors.png')
fig,axs=plt.subplots(1,2,figsize=(9,4.2));counts=[24,13,3];labels=['Matched in\n0 ≤ e/D ≤ 0.05','Target below\ncorrected e/D=0','Target above\ne/D=0.05']
axs[0].bar(labels,counts,color=['#'+TEAL,'#'+ORANGE,'#'+RED]);axs[0].set_ylim(0,28);axs[0].set_ylabel('Dimpled cases');axs[0].tick_params(axis='x',labelsize=9)
for i,v in enumerate(counts):axs[0].text(i,v+.5,str(v),ha='center',weight='bold')
merged=dimple.merge(old[['Case','Fitted effective e/D']],on='Case');matched=merged.dropna(subset=['effective_e_over_D'])
axs[1].scatter(matched['Fitted effective e/D'],matched.effective_e_over_D,color='#'+TEAL,s=30)
axs[1].axhline(.05,color='#'+ORANGE,ls='--');axs[1].set(xlabel='Historical fitted e/D',ylabel='New bounded fitted e/D',title='Matched subset only (24 cases)');fig.tight_layout();save(fig,'dimple_transfer.png')
fig,ax=plt.subplots(figsize=(8.5,4));cases=['May19_3','May27_10'];sample=dimple.set_index('Case').loc[cases];x=np.arange(2)
a=ax.bar(x-.17,sample.heat_flux_CO2_W_m2,.34,color='#'+TEAL,label='Numerical, HTC fitted');b=ax.bar(x+.17,sample.nominal_heat_flux_W_m2,.34,color='#'+ORANGE,label='Nominal experimental group')
ax.bar_label(a,fmt='%.0f',padding=4);ax.bar_label(b,fmt='%.0f',padding=4);ax.set_xticks(x,cases);ax.set_ylim(0,3500);ax.set_ylabel('Heat flux (W/m²)');ax.legend(fontsize=9,loc='upper left');save(fig,'flux_examples.png')
fig,ax=plt.subplots(figsize=(8.8,3.8));x=np.linspace(0,1,301)
for name,n,c in [('Constant',1,TEAL),('Step',3,BLUE),('Linear',1,ORANGE),('Sine',1,RED)]:ax.plot(x,build_roughness(x,.02,name,n),label=name+(' (3 steps)' if name=='Step' else ''),color='#'+c,lw=2)
ax.set(xlabel='Axial position x/L',ylabel='Local e/D',ylim=(-.001,.022));ax.legend(ncol=4,fontsize=9,loc='upper left');save(fig,'roughness_profiles.png')
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def rgb(v):return RGBColor.from_string(v)
def textbox(s,text,x,y,w,h,size=20,color=NAVY,bold=False):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold;p.space_after=Pt(11)
 return box
slidesources=[]
def slide(title,sub='',source='Repository results; code snapshot da28165',notes=''):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FFFFFF')
 sh=s.shapes.add_shape(1,0,0,prs.slide_width,Inches(.13));sh.fill.solid();sh.fill.fore_color.rgb=rgb(TEAL);sh.line.fill.background()
 textbox(s,title,.55,.38,12.2,.62,29,bold=True)
 if sub:textbox(s,sub,.57,1.06,12.1,.64,14,GRAY)
 textbox(s,source,.57,7.08,11.5,.25,9,GRAY);textbox(s,str(len(prs.slides)),12.12,7.03,.65,.3,11,GRAY)
 s.notes_slide.notes_text_frame.text=notes+'\nSources: '+source+'\nRepository: https://github.com/QuinAcchione52/counterflow-hx-model\nModel-results snapshot: da28165; prepared October 13, 2026.'
 return s

def bullets(s,items,x=.65,y=1.85,w=11.9,size=22):
 text='\n'.join('• '+t for t in items);return textbox(s,text,x,y,w,4.8,size)
def callout(s,text,y=6.35,color=TEAL):textbox(s,text,.65,y,12,.55,17,color,bold=True)
def picture(s,path,x=.65,y=1.8,w=8.3):
 with Image.open(path) as im: ratio=im.width/im.height
 actual=min(w,4.35*ratio)
 s.shapes.add_picture(str(path),Inches(x+(w-actual)/2),Inches(y),width=Inches(actual))
def table(s,headers,rows,x=.65,y=1.9,w=12,h=3.8,size=15,widths=None):
 shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h));t=shape.table
 if widths:
  for col,ww in zip(t.columns,widths):col.width=Inches(ww)
 for i,row in enumerate([headers]+rows):
  for j,val in enumerate(row):
   cell=t.cell(i,j);cell.text=str(val);cell.margin_left=Inches(.09);cell.margin_right=Inches(.06)
   cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if i==0 else ('F1F5F8' if i%2 else 'FFFFFF'))
   for p in cell.text_frame.paragraphs:p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=(i==0);p.font.color.rgb=rgb('FFFFFF' if i==0 else NAVY)
 return t

s=slide('Counterflow sCO₂ heat-exchanger model','Code audit • smooth baseline testing • dimpled HTC calibration',source='Spencer thesis + technical note; repository snapshot da28165')
textbox(s,'What changed, what improved,\nand what remains unresolved',.75,2.15,11.6,1.85,36,bold=True)
textbox(s,'38 smooth cases  |  40 dimpled cases  |  HEOS properties',.8,4.2,11.7,.65,24,TEAL)
textbox(s,'Research review • October 13, 2026\nAll e/D matches are conditional on stated operating assumptions.',.8,5.35,11.5,.9,18,GRAY)
s=slide('The main result: improvement with clear limits')
table(s,['Question','Result'],[['Did real code errors exist?','Yes: profiles, property handling, friction and local wall closure.'],['Did those fixes explain the HTC gap?','No: representative smooth friction change is only about 1%.'],['Did smooth calibration help prediction?','Held-out mean absolute HTC error: 46.73% → 14.22%.'],['Can bounded e/D match more dimples?','24/40 matched within 0–0.05; 16 remain unfitted.'],['Is the physical mismatch fully solved?','No: heat flux, boundary conditions and applicability remain issues.']],h=4.1,size=18,widths=[3.7,8.3])
callout(s,'A calibrated HTC match is not independent validation of roughness, duty or pressure drop.')
s=slide('Starting problem and the analysis sequence','Historical observations are not a controlled before/after comparison.')
bullets(s,['Initial May27 result: roughly 543 versus experimental HTC around 2600 W/(m²·K).','Changing guessed inlet temperature and allowing very large e/D could fit HTC, but interpretation was unclear.','First compare smooth tubes at e/D=0; then audit the code; then test a separate smooth baseline correction.','Hold the smooth-trained correction fixed while fitting individual dimpled e/D.'])
callout(s,'The experimental observations were not edited to improve agreement.')
s=slide('What data were used—and what was assumed',source='Spencer thesis Tables 12–13, printed pp.55–56; Smooth_cases.csv; Dimpled_cases.csv')
table(s,['Input / target','Treatment'],[['38 smooth + 40 dimpled cases','Pressure, CO₂ flow, mean temperature and HTC from source tables.'],['CO₂ inlet / outlet temperatures','Inlet inferred to reproduce the reported endpoint mean; not measured inputs.'],['Water inlet / flow','Shared assumptions: 20.7 °C and 0.087 kg/s.'],['Geometry / properties','ID 13.3858 mm; length 1 m; wall 1.2446 mm; HEOS backend.'],['Nominal heat flux / uncertainty','2 or 3 kW/m² group; retain source HTC uncertainty unchanged.']],h=4.1,size=17,widths=[3.8,8.2])
callout(s,'Nominal flux group ≠ exact measured per-case duty; constant water conditions ≠ exact rig reproduction.')
s=slide('Compare the same HTC definition','The endpoint metric and the mean of local HTC values answer different questions.',source='hx_metrics.py; CounterHX_full.py; Spencer experimental definition')
bullets(s,['Reported comparison metric: h = Q / (Aᵢ × ΔTₗₘ,bulk–inner wall).','This LMTD is the CO₂ bulk-to-wall difference at the two tube ends—not CO₂-to-water LMTD.','The CO₂ inlet is solved jointly so (Tᵢₙ + Tₒᵤₜ)/2 matches the reported mean.','The output also preserves local HTC, wall temperatures, duty and energy residuals.'])
callout(s,'For nonuniform e/D, endpoint HTC and length-mean local HTC can differ substantially.')
s=slide('Code changes 1/3: roughness and DH baseline',source='roughness.py; Correlations.py; validation/correlation_audit.md')
table(s,['Before','After'],[['Step with n=1 produced zero e/D','One step applies the requested constant e/D.'],['Custom arrays reduced to first value','Each axial node retains its supplied roughness value.'],['Step boundaries depended on node partition','Profiles use normalized physical position x/L.'],['Direct Colebrook replaced smooth DH friction','Original smooth DH recovered exactly at e/D=0.'],['Ambiguous rough-tube DH interpretation','Relative Colebrook enhancement explicitly labeled empirical.']],h=4.1,size=17,widths=[5.3,6.7])
s=slide('Local roughness profiles now behave as requested',source='roughness.py; illustration: maximum e/D=0.02')
picture(s,ASSETS/'roughness_profiles.png',w=9.1)
textbox(s,'Validated behavior\n\nOne Step = Constant.\n\nCustom arrays preserved.\n\nNegative / malformed inputs rejected.',10,2.1,2.6,3.7,18)
callout(s,'Profile comparison is illustrative; not an experimentally validated dimple geometry model.')
s=slide('Code changes 2/3: properties and wall iteration',source='Correlations.py; Counterflow_HX_4.py')
bullets(s,['Removed silent clipping of CO₂ and wall temperatures; invalid states produce explicit errors.','Water pressure is supplied in Pa to the correlation; no unit guessing based on magnitude.','Water-wall iteration now converges to a tolerance instead of stopping after three iterations.','Coupled solver starts at zero duty rather than an arbitrary 3 kW/m² seed that could create a false freezing state.','CoolProp named phase constants replace incorrect numerical labels.'],size=21)
s=slide('Code changes 3/3: balances and pressure bookkeeping',source='Counterflow_HX_4.py; tests/test_correlations_profiles.py')
bullets(s,['Node heat flux now follows local conductance and bulk temperature difference. Segment duty remains integrated separately.','Convergence checks include local water-wall heat-flux closure; both wall-flux and conduction residuals are exported.','Thermal friction uses film Reynolds number; mechanical Darcy friction uses bulk Reynolds number.','Added diagnostics: selected Prandtl branch, Cp averaging, film conductivity, phase and separate friction factors.'],size=21)
callout(s,'Corrected pressure bookkeeping is not an experimental pressure-drop calibration.')
s=slide('Dang–Hihara remains the baseline','Equation audit uses Spencer’s transcription; original article was not independently verified.',source='Spencer thesis Eqs.4–10, printed pp.9–10; Correlations.py')
textbox(s,'Smooth thermal friction\nfDH = [1.82 log₁₀(Refilm) − 1.64]⁻²',.8,1.85,11.6,1.15,26,bold=True)
textbox(s,'Local roughness extension\nfheat = fDH × fColebrook(Refilm, e/D) / fColebrook(Refilm, 0)',.8,3.15,11.6,1.15,23,TEAL,bold=True)
textbox(s,'Optional empirical baseline\nNucorrected = Csmooth × NuDH,roughness',.8,4.55,11.6,1,24,bold=True)
callout(s,'The roughness ratio and smooth multiplier are empirical additions; defaults keep Csmooth = 1.')
s=slide('Fixing the smooth friction expression is a small change',source='validation/Correlation_audit_first_pass_HEOS.csv; 40 segments, shared water assumptions')
rows=[]
for case,ep in [('May27_10',0),('May27_10',.05),('May19_3',0)]:
 a=audit[(audit.Case==case)&(audit['e/D']==ep)];before=float(a[a.model=='colebrook'].iloc[0].htc_spencer_W_m2K);after=float(a[a.model=='dh_roughness'].iloc[0].htc_spencer_W_m2K)
 rows.append([case,f'{ep:.2f}',f'{before:.2f}',f'{after:.2f}',f'{100*(after/before-1):+.2f}%'])
table(s,['Case','e/D','Direct Colebrook HTC','DH-based HTC','Change'],rows,h=2.4,size=18,widths=[2.3,1,3,3,2.7])
textbox(s,'Units: W/(m²·K). This isolates thermal friction choices within the corrected solver; it is not a replay of the initial 543 result.',.8,4.65,11.6,1,20,GRAY)
callout(s,'The large HTC gap cannot be attributed to this friction-expression difference alone.')
s=slide('Why a baseline mismatch can persist',source='Spencer thesis §7.2 / Table 8, printed pp.38–39; procedure printed p.29; correlation_audit.md')
bullets(s,['The cited DH database uses smaller smooth tubes (1–6 mm) and substantially higher mass flux (~200 kg/m²·s).','Spencer’s modeled tube is 13.4 mm ID; smooth-case mass flux is about 21–35 kg/m²·s.','Buoyancy, thermal entrance effects and dimple-induced secondary flows are not independently modeled by the roughness ratio.','The source procedure adjusts or prechills water; shared 20.7 °C water inputs are an approximation.','Smooth-tube mismatch exists before adding dimples. This does not establish experimental error.'],size=21)
s=slide('How the optional smooth correction was tested',source='train_smooth_baseline.py; smooth_baseline.py; model.json')
bullets(s,['Fit one positive multiplier per smooth case at e/D=0 inside the coupled solver.','Fit log(multiplier) using temperature relative to the pseudocritical point, its square, log(mass flux) and log(nominal flux group).','For each held-out case, exclude its fitted factor; train on the other 37; predict its multiplier and run the HX solver.','Then train on all 38 smooth cases for transfer to dimples. No dimple HTC is used to train the smooth baseline.'],size=21)
callout(s,'Nominal flux is a predictor, not imposed heat duty. Related campaign cases limit generalization claims.')
s=slide('Smooth predictions improve, but are not all within 5%',source='results/smooth_baseline/heldout_predictions.csv; 38 cases, 20 segments')
picture(s,ASSETS/'smooth_parity.png',w=8.4)
textbox(s,'Mean absolute error\n46.73% → 14.22%\n\nWithin ±5%\n9 / 38 cases\n\nWorst held-out error\n43.46%',9.35,2.1,3.2,3.9,22,TEAL,bold=True)
s=slide('All 38 smooth cases: signed error comparison',source='results/smooth_baseline/heldout_predictions.csv; shaded region: ±5% target error')
picture(s,ASSETS/'smooth_case_errors.png',x=.75,y=1.8,w=11.85)
callout(s,'Percentage prediction error is distinct from Spencer’s reported experimental uncertainty.',y=6.53)
s=slide('Dimpled transfer: 24 matches within the tested e/D bound',source='Dimpled-transfer.summary.csv; historical all-40-case CSV; 20-segment screen')
picture(s,ASSETS/'dimple_transfer.png',w=10.6)
callout(s,'Historical: 4/40 within 0–0.05. Current: 24/40. Different formulation; neither count proves physical roughness.')
s=slide('Two example fits confirmed at 100 segments',source='Dimpled_fine_grid_checks.csv; historical all-40-case CSV')
rows=[]
for _,r in fine.iterrows():
 prior=old[old.Case==r.Case].iloc[0]
 rows.append([r.Case,f'{prior["Fitted effective e/D"]:.6f}',f'{r.effective_e_over_D:.7f}',f'{r.target_htc_W_m2K:.0f}',f'{r.htc_spencer_W_m2K:.2f}'])
table(s,['Case','Historical e/D¹','New e/D','Target HTC','Numerical HTC'],rows,h=2.1,size=20,widths=[2.4,2.4,2.4,2.4,2.4])
textbox(s,'¹ Historical 20-segment calibrated table; current example fits use 100 segments.\nHTC units: W/(m²·K). These are fitted targets, not held-out dimple predictions.',.8,4.35,11.6,1.15,18,GRAY)
callout(s,'A smaller fitted e/D reflects the added baseline correction; it is not an independent roughness measurement.')
s=slide('Matching HTC still does not match heat flux',source='Dimpled-transfer.summary.csv; matched 20-segment solutions')
picture(s,ASSETS/'flux_examples.png',w=8.6)
textbox(s,'May19_3\n1436 vs 2000 W/m²\n≈ −28.2%\n\nMay27_10\n1412 vs 3000 W/m²\n≈ −52.9%',9.6,2.2,3,3.65,21,ORANGE,bold=True)
callout(s,'These are differences from nominal groups; exact measured per-case duty is needed for a definitive comparison.')
s=slide('How the data workflow improved')
bullets(s,['Experimental target HTC, reported uncertainty and nominal flux remain separate from numerical outputs.','Measured CO₂ pressure/flow/mean temperature are distinguished from assumed water conditions and inferred inlet temperature.','All-case outputs preserve unfitted status, baseline extrapolation flags and local diagnostics.','Historical tables remain archived; new formulation results are saved separately.','Prediction errors and fitted target agreement are labeled separately. No experimental uncertainty was reduced.'],size=21)
callout(s,'Better traceability and prediction—not modification of the experimental data.')
s=slide('Verification completed—and its practical limits',source='tests/test_correlations_profiles.py; Smooth_grid_checks.csv; Dimpled_fine_grid_checks.csv')
bullets(s,['17 tests pass: profile semantics, Colebrook reference, DH smooth limit, invalid-state guards, wall and energy balances, and calibration-domain checks.','Three representative corrected smooth cases: 20 → 100 segments changed HTC by less than 0.002%.','May19_3 and May27_10 dimple fits confirmed at 100 segments.','Held-out smooth energy residuals are below 6 × 10⁻⁸ W; fitted dimple CO₂ residuals below 3.3 × 10⁻⁶ W.','May24_1 is flagged outside a smooth-training range. No pressure-drop agreement is established.'],size=20)
callout(s,'Numerical convergence and energy closure do not establish physical validity.')
s=slide('Recommended discussion with your professor')
bullets(s,['“We corrected identifiable implementation errors while retaining Dang–Hihara.”','“Those fixes do not explain the full gap; a separate smooth-trained correction improves held-out HTC predictions.”','“Bounded dimple fits increase to 24/40, but these are empirical calibrations and heat flux still disagrees.”','Next: recover per-case inlet/outlet, water flow/temperature, measured duty and wall-temperature data.','Then test entrance / mixed-convection physics and validate on an independent campaign before interpreting e/D as roughness.'],size=21)
callout(s,'The evidence supports model-development progress, not a claim that the rig or correlation has been disproved.')
# Complete per-case appendices, 10 rows per slide.
for start in range(0,len(dimple),10):
 chunk=dimple.iloc[start:start+10]
 s=slide(f'Appendix: dimpled cases {start+1}–{start+len(chunk)} of 40','HTC in W/(m²·K). Unfitted cases show endpoint predictions; no forced match.',source='results/smooth_baseline/Dimpled-transfer.summary.csv; 20 segments')
 rows=[]
 for _,r in chunk.iterrows():
  fit=pd.notna(r.effective_e_over_D);status='Matched' if fit else ('Below baseline' if r.target_htc_W_m2K<r.smooth_model_htc_W_m2K else 'Above cap')
  if r.baseline_model_extrapolated:status+=' *'
  rows.append([r.Case,f'{r.target_htc_W_m2K:.0f}',f'{r.smooth_model_htc_W_m2K:.0f}',f'{r.htc_at_eD_005_W_m2K:.0f}',f'{r.htc_spencer_W_m2K:.0f}' if fit else '—',f'{r.effective_e_over_D:.6f}' if fit else '—',status])
 table(s,['Case','Target','e/D=0 HTC','e/D=.05 HTC','Fitted HTC','Fitted e/D','Status'],rows,h=4.35,size=13,widths=[2.15,1.2,1.7,1.8,1.55,1.45,2.15])
 callout(s,'* May24_1: smooth-baseline range extrapolation. A tested e/D bound is not a validated physical limit.',y=6.45)
for start in range(0,len(smooth),10):
 chunk=smooth.iloc[start:start+10]
 s=slide(f'Appendix: smooth cases {start+1}–{start+len(chunk)} of 38','HTC in W/(m²·K); correction predicted with this case withheld; e/D=0.',source='results/smooth_baseline/heldout_predictions.csv; 20 segments')
 rows=[]
 for _,r in chunk.iterrows():rows.append([r.Case,f'{r.target_htc_W_m2K:.0f}',f'{r.uncorrected_htc_W_m2K:.1f}',f'{r.heldout_numerical_htc_W_m2K:.1f}',f'{r.heldout_error_percent:+.1f}%', 'Yes' if abs(r.heldout_error_percent)<=5 else 'No'])
 table(s,['Case','Target','Uncorrected','Held-out HTC','Error','Within 5%'],rows,h=4.35,size=14,widths=[2.6,1.6,2.1,2.2,1.7,1.8])
s=slide('Sources, reproducibility and files',source='GitHub: QuinAcchione52/counterflow-hx-model; model-results snapshot da28165')
bullets(s,['Spencer thesis: Eqs.4–10 (pp.9–10), procedure (p.29), correlation comparison (pp.38–39), Tables 12–13 (pp.55–56); supplied technical note as companion reference.','Code audit: validation/correlation_audit.md. Implementation: Correlations.py, Counterflow_HX_4.py, roughness.py.','Results: results/smooth_baseline/heldout_predictions.xlsx and Dimpled-transfer.xlsx.','Reproduce correction: train_smooth_baseline.py. Transfer fitting: calibrate_dimpled.py --smooth-baseline-model results/smooth_baseline/model.json --ep-max .05.','See repository README for full Windows commands, explicit baseline-extrapolation flag and HEOS requirements.'],size=18)
callout(s,'Primary Dang–Hihara article not independently verified; equation audit relies on Spencer’s transcription.')
prs.core_properties.title='Counterflow sCO2 HX — code audit and calibration results'
prs.core_properties.subject='Dang–Hihara baseline, held-out smooth correction, and bounded dimpled transfer'
prs.core_properties.author='Counterflow HX model research review'
prs.core_properties.comments='Generated from repository snapshot da28165; experimental data not modified.'
path=OUT/'Counterflow_HX_Code_and_Results_Review.pptx';prs.save(path)
print(f'Saved {len(prs.slides)} slides: {path}')
